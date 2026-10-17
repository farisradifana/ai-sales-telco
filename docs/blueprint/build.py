"""Render the blueprint (content.py) to HTML, TXT and DOCX with identical content.

  python docs/blueprint/build.py            # needs python-docx for the .docx (pip install python-docx)
"""
import html
import sys
import textwrap
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from content import META, SECTIONS, SUBTITLE, TITLE  # noqa: E402

OUT = Path(__file__).resolve().parent
BASE = 'AI_Sales_Intelligence_Telco_Blueprint'


# ---------------------------------------------------------------- HTML
def render_html() -> str:
    e = html.escape
    toc = ''.join(f'<li><a href="#{sid}">{e(title)}</a></li>' for sid, title, _ in SECTIONS)
    meta = ''.join(f'<div class="kv"><span>{e(k)}</span><b>{e(v)}</b></div>' for k, v in META)
    body = []
    for sid, title, blocks in SECTIONS:
        body.append(f'<section id="{sid}"><h2>{e(title)}</h2>')
        for b in blocks:
            kind = b[0]
            if kind == 'p':
                body.append(f'<p>{e(b[1])}</p>')
            elif kind == 'h3':
                body.append(f'<h3>{e(b[1])}</h3>')
            elif kind in ('ul', 'ol'):
                body.append(f'<{kind}>' + ''.join(f'<li>{e(x)}</li>' for x in b[1]) + f'</{kind}>')
            elif kind == 'code':
                body.append(f'<pre>{e(b[1])}</pre>')
            elif kind == 'table':
                head = ''.join(f'<th>{e(h)}</th>' for h in b[1])
                rows = ''.join('<tr>' + ''.join(f'<td>{e(c)}</td>' for c in r) + '</tr>' for r in b[2])
                body.append(f'<div class="tw"><table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table></div>')
        body.append('</section>')
    return f'''<!doctype html>
<html lang="id"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{e(TITLE)} — Blueprint</title>
<style>
:root{{--bg:#f7f8fb;--card:#ffffff;--ink:#111827;--muted:#4b5563;--line:#e5e7eb;--accent:#2563eb;--code:#0f172a;--codeink:#e2e8f0;--head:#eef2ff}}
@media (prefers-color-scheme:dark){{:root{{--bg:#0b1020;--card:#111827;--ink:#e5e7eb;--muted:#9ca3af;--line:#1f2937;--accent:#60a5fa;--code:#020617;--codeink:#e2e8f0;--head:#1e293b}}}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--ink);font:16px/1.6 system-ui,Segoe UI,Roboto,sans-serif}}
header{{background:linear-gradient(135deg,#1d4ed8,#7c3aed);color:#fff;padding:48px 16px}}
header .in,main{{max-width:1080px;margin:0 auto}}header h1{{margin:0 0 6px;font-size:clamp(26px,5vw,40px)}}header p{{margin:0;opacity:.9}}
.meta{{display:grid;grid-template-columns:repeat(auto-fit,minmax(240px,1fr));gap:10px;margin-top:22px}}
.kv{{background:rgba(255,255,255,.12);border-radius:10px;padding:10px 12px}}.kv span{{display:block;font-size:12px;opacity:.85}}.kv b{{font-weight:600}}
main{{padding:24px 16px 64px;display:grid;grid-template-columns:240px 1fr;gap:28px}}
nav{{position:sticky;top:16px;align-self:start}}nav ol{{list-style:none;margin:0;padding:0}}nav a{{display:block;padding:6px 10px;border-radius:8px;color:var(--muted);text-decoration:none}}nav a:hover{{background:var(--card);color:var(--accent)}}
section{{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:22px 24px;margin-bottom:18px}}
h2{{margin:0 0 12px;font-size:22px}}h3{{margin:20px 0 8px;font-size:17px;color:var(--accent)}}
.tw{{overflow-x:auto}}table{{border-collapse:collapse;width:100%;font-size:14px;margin:8px 0}}th,td{{border:1px solid var(--line);padding:8px 10px;text-align:left;vertical-align:top}}th{{background:var(--head)}}
pre{{background:var(--code);color:var(--codeink);padding:14px 16px;border-radius:10px;overflow-x:auto;font:13px/1.5 ui-monospace,Consolas,monospace;white-space:pre}}
li{{margin:4px 0}}footer{{text-align:center;color:var(--muted);font-size:13px;padding:0 16px 32px}}
@media (max-width:820px){{main{{grid-template-columns:1fr}}nav{{position:static}}section{{padding:18px 16px}}}}
@media print{{nav{{display:none}}main{{display:block}}header{{-webkit-print-color-adjust:exact;print-color-adjust:exact}}}}
</style></head><body>
<header><div class="in"><h1>{e(TITLE)}</h1><p>{e(SUBTITLE)}</p><div class="meta">{meta}</div></div></header>
<main><nav aria-label="Daftar isi"><ol>{toc}</ol></nav><div>{''.join(body)}</div></main>
<footer>{e(TITLE)} · blueprint dibuat dari docs/blueprint/content.py</footer>
</body></html>
'''


# ---------------------------------------------------------------- TXT
def _wrap(text, width=100, indent=''):
    return textwrap.fill(text, width=width, initial_indent=indent, subsequent_indent=' ' * len(indent)) if text else ''


def _txt_table(headers, rows, width=100):
    n = len(headers)
    natural = [max(len(str(r[i])) for r in [headers] + rows) for i in range(n)]
    avail = width - (3 * n + 1)
    cols = [min(w, max(12, avail // n)) for w in natural]
    spare = avail - sum(cols)
    for i in sorted(range(n), key=lambda i: natural[i] - cols[i], reverse=True):
        if spare <= 0:
            break
        add = min(spare, natural[i] - cols[i])
        if add > 0:
            cols[i] += add
            spare -= add
    sep = '+' + '+'.join('-' * (c + 2) for c in cols) + '+'

    def line(cells):
        wrapped = [textwrap.wrap(str(c), cols[i]) or [''] for i, c in enumerate(cells)]
        out = []
        for k in range(max(len(w) for w in wrapped)):
            out.append('| ' + ' | '.join((w[k] if k < len(w) else '').ljust(cols[i]) for i, w in enumerate(wrapped)) + ' |')
        return out
    return [sep, *line(headers), sep.replace('-', '='), *[x for r in rows for x in line(r) + [sep]]]


def render_txt() -> str:
    out = ['=' * 100, TITLE.upper().center(100), SUBTITLE.center(100), '=' * 100, '']
    out += [f'{k:<16}: {v}' for k, v in META] + ['', 'DAFTAR ISI'] + [f'  {t}' for _, t, _ in SECTIONS] + ['']
    for _, title, blocks in SECTIONS:
        out += ['', title.upper(), '-' * len(title), '']
        for b in blocks:
            kind = b[0]
            if kind == 'p':
                out += [_wrap(b[1]), '']
            elif kind == 'h3':
                out += [f'## {b[1]}', '']
            elif kind == 'ul':
                out += [_wrap(x, indent='  - ') for x in b[1]] + ['']
            elif kind == 'ol':
                out += [_wrap(x, indent=f'  {i}. ') for i, x in enumerate(b[1], 1)] + ['']
            elif kind == 'code':
                out += ['    ' + ln for ln in b[1].splitlines()] + ['']
            elif kind == 'table':
                out += _txt_table(b[1], b[2]) + ['']
    return '\n'.join(out).rstrip() + '\n'


# ---------------------------------------------------------------- DOCX
def render_docx(path: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(2)
    st = doc.styles['Normal']
    st.font.name, st.font.size = 'Calibri', Pt(10.5)

    def shade(cell, hex_fill):
        tcpr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement('w:shd')
        sh.set(qn('w:val'), 'clear')
        sh.set(qn('w:color'), 'auto')
        sh.set(qn('w:fill'), hex_fill)
        tcpr.append(sh)

    t = doc.add_heading(TITLE, level=0)
    doc.add_paragraph(SUBTITLE).runs[0].italic = True
    mt = doc.add_table(rows=0, cols=2)
    mt.style = 'Table Grid'
    for k, v in META:
        r = mt.add_row().cells
        r[0].text, r[1].text = k, v
        r[0].paragraphs[0].runs[0].bold = True
        shade(r[0], 'EEF2FF')
    doc.add_paragraph()
    doc.add_heading('Daftar isi', level=1)
    for _, title, _ in SECTIONS:
        doc.add_paragraph(title, style='List Bullet')
    for _, title, blocks in SECTIONS:
        doc.add_page_break() if title.startswith(('3.', '5.', '7.', '9.')) else None
        doc.add_heading(title, level=1)
        for b in blocks:
            kind = b[0]
            if kind == 'p':
                doc.add_paragraph(b[1])
            elif kind == 'h3':
                doc.add_heading(b[1], level=2)
            elif kind == 'ul':
                for x in b[1]:
                    doc.add_paragraph(x, style='List Bullet')
            elif kind == 'ol':
                for x in b[1]:
                    doc.add_paragraph(x, style='List Number')
            elif kind == 'code':
                p = doc.add_paragraph()
                run = p.add_run(b[1])
                run.font.name, run.font.size = 'Consolas', Pt(8.5)
                run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
                p.paragraph_format.left_indent = Cm(0.5)
            elif kind == 'table':
                tb = doc.add_table(rows=1, cols=len(b[1]))
                tb.style = 'Table Grid'
                tb.alignment = WD_TABLE_ALIGNMENT.CENTER
                for i, h in enumerate(b[1]):
                    c = tb.rows[0].cells[i]
                    c.text = h
                    c.paragraphs[0].runs[0].bold = True
                    shade(c, 'DBEAFE')
                for row in b[2]:
                    cells = tb.add_row().cells
                    for i, v in enumerate(row):
                        cells[i].text = str(v)
                for row in tb.rows:
                    for c in row.cells:
                        for p in c.paragraphs:
                            for r in p.runs:
                                r.font.size = Pt(9)
                doc.add_paragraph()
    doc.save(path)


def main():
    (OUT / f'{BASE}.html').write_bytes(render_html().encode('utf-8'))
    (OUT / f'{BASE}.txt').write_bytes(render_txt().encode('utf-8'))
    render_docx(OUT / f'{BASE}.docx')
    for ext in ('html', 'txt', 'docx'):
        f = OUT / f'{BASE}.{ext}'
        print(f.name, f.stat().st_size, 'bytes')


if __name__ == '__main__':
    main()
